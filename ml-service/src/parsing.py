"""
Resume document parsing: PDF, DOCX and plain text → normalised text.

Kept dependency-light (pypdf + python-docx) so the service installs in seconds
without a GPU toolchain.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from typing import Optional


class ParseError(Exception):
    """Raised when a document cannot be read at all."""


@dataclass
class ParsedDocument:
    text: str
    pages: Optional[int]
    engine: str


BULLET_CHARS = "•●▪◦‣·"
CONTROL_CHARS = re.compile(r"[\u0000-\u0008\u000b\u000c\u000e-\u001f]")


def normalize_text(raw: str) -> str:
    """Collapse PDF extraction artefacts into readable text."""
    text = raw.replace("\r\n", "\n").replace("\r", "\n")
    text = CONTROL_CHARS.sub(" ", text)
    text = text.replace("\u00a0", " ")
    for bullet in BULLET_CHARS:
        text = text.replace(bullet, "- ")
    text = re.sub(r"[ \t]{2,}", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    # Re-join words hyphenated across line breaks ("engineer-\ning").
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
    # Bullet glyphs sometimes arrive as lone dashes on their own line.
    text = re.sub(r"\n\s*[-–—]\s*\n", "\n- ", text)
    return text.strip()


def parse_pdf(payload: bytes) -> ParsedDocument:
    try:
        from pypdf import PdfReader  # type: ignore
    except ImportError as error:  # pragma: no cover - depends on environment
        raise ParseError("pypdf is not installed on the ML service.") from error

    try:
        reader = PdfReader(io.BytesIO(payload))
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as error:
        raise ParseError(f"Could not read that PDF: {error}") from error

    text = normalize_text("\n\n".join(pages))
    if len(text) < 40:
        raise ParseError(
            "This PDF has no extractable text layer (it looks like a scan). "
            "Export a text-based PDF, or paste the resume text instead."
        )
    return ParsedDocument(text=text, pages=len(pages), engine="pypdf")


def parse_docx(payload: bytes) -> ParsedDocument:
    try:
        import docx  # type: ignore
    except ImportError as error:  # pragma: no cover - depends on environment
        raise ParseError("python-docx is not installed on the ML service.") from error

    try:
        document = docx.Document(io.BytesIO(payload))
    except Exception as error:
        raise ParseError(f"Could not read that DOCX: {error}") from error

    parts: list[str] = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return ParsedDocument(text=normalize_text("\n".join(parts)), pages=None, engine="python-docx")


def parse_document(filename: str, payload: bytes) -> ParsedDocument:
    lower = filename.lower()
    if lower.endswith(".pdf") or payload[:4] == b"%PDF":
        return parse_pdf(payload)
    if lower.endswith(".docx") or payload[:2] == b"PK":
        return parse_docx(payload)
    if lower.endswith((".txt", ".md", ".markdown", ".rst", ".json")):
        try:
            return ParsedDocument(text=normalize_text(payload.decode("utf-8", errors="ignore")), pages=None, engine="text")
        except Exception as error:
            raise ParseError(f"Could not read that text file: {error}") from error

    # Last resort: try to decode it as text before giving up.
    try:
        decoded = payload.decode("utf-8")
        if decoded.count("\ufffd") < 10 and len(decoded.strip()) > 200:
            return ParsedDocument(text=normalize_text(decoded), pages=None, engine="text-fallback")
    except UnicodeDecodeError:
        pass

    raise ParseError(
        "Unsupported file type. Upload a PDF, DOCX or TXT file — or paste the resume text into ApplyAI."
    )
